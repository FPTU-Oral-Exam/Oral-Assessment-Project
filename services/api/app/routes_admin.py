import hashlib
import random
import time
from copy import deepcopy
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, UploadFile
from fastapi.responses import Response
from sqlalchemy import delete, func, or_, select
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from . import ai, storage
from . import schemas as s
from .course_deletion import delete_course_tree
from .db import get_db
from .grading import GradingError, assessment_view, check_config, review_question
from .knowledge import chunk_scope, set_mappings, topic_data
from .models import (
    Assignment,
    Attempt,
    Audit,
    BookSection,
    Course,
    CourseEnrollment,
    Document,
    Exam,
    ExamBatch,
    ExamEnrollment,
    ExamSession,
    LearningOutcome,
    MediaCleanup,
    QuestionItem,
    ReEvaluation,
    ReviewJob,
    Rubric,
    Section,
    Topic,
    TopicDocument,
    Upload,
    User,
    uid,
)
from .retakes import allowance, history_row, sessions_for
from .runtime_settings import settings
from .security import admin, by_id, course_access, editor, examiner, fail, hasher, public_user, staff
from .speech import google_ready, policy
from .worker import finalize, grade_answer

router = APIRouter()


def data(row, *fields):
    return {k: getattr(row, k) for k in ("id", *fields)}


def course_list(db, user):
    query = select(Course).order_by(Course.created_at.desc())
    return db.scalars(query).all()


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), user=Depends(staff)):
    ids = [c.id for c in course_list(db, user)]
    return {
        "courses": len(ids),
        "exams": db.scalar(select(func.count()).select_from(Exam).where(Exam.course_id.in_(ids))),
        "items": db.scalar(select(func.count()).select_from(QuestionItem).where(QuestionItem.course_id.in_(ids), QuestionItem.deleted_at.is_(None))) or 0,
        "sessions": db.scalar(
            select(func.count()).select_from(ExamSession).join(Exam).where(Exam.course_id.in_(ids), ExamSession.deleted_at.is_(None))
        ),
        "ai_provider": settings().ai_provider,
    }


@router.get("/users")
def users(role: str | None = None, db: Session = Depends(get_db), user=Depends(staff)):
    query = select(User).order_by(User.created_at.desc())
    if role:
        query = query.where(User.role == role)
    elif user.role not in {"SYSTEM_ADMIN", "ADMIN", "EXAMINER", "REVIEWER"}:
        query = query.where(User.role == "STUDENT")
    return [public_user(u) for u in db.scalars(query)]


@router.post("/users", status_code=201)
def create_user(body: s.UserIn, db: Session = Depends(get_db), user=Depends(admin)):
    row = User(**body.model_dump(exclude={"password"}), password_hash=hasher.hash(body.password))
    db.add(row)
    db.add(Audit(user_id=user.id, event="USER_CREATED", details={"username": row.username}))
    db.commit()
    return public_user(row)


@router.get("/courses")
def courses(db: Session = Depends(get_db), user=Depends(staff)):
    return [data(c, "code", "name", "description", "status", "owner_id", "teacher_id") for c in course_list(db, user)]


@router.get("/teacher-semesters")
def teacher_semesters(db: Session = Depends(get_db), user=Depends(staff)):
    """Return all semesters that have batches assigned to the current teacher,
    or courses where teacher is assigned/owner.
    Each semester entry includes the courses (with exams and batch info) assigned to this teacher.
    Used by the Teacher Portal landing page to show a semester-first navigation.
    """
    from .models import ExamBatch, Semester, ScheduleSlot

    is_super = user.role in {"SYSTEM_ADMIN", "ADMIN", "EXAMINER"}

    # 1. Batches assigned to teacher (or all if superuser)
    if is_super:
        batches = db.scalars(select(ExamBatch)).all()
    else:
        batches = db.scalars(
            select(ExamBatch).where(ExamBatch.assigned_teacher_id == user.id)
        ).all()

    # 2. Courses directly assigned to teacher (or owned by teacher)
    if is_super:
        direct_courses = db.scalars(select(Course)).all()
    else:
        direct_courses = db.scalars(
            select(Course).where(or_(Course.teacher_id == user.id, Course.owner_id == user.id))
        ).all()

    # Pre-load exams and courses for batches
    exam_ids = list({b.exam_id for b in batches})
    batch_exams = db.scalars(select(Exam).where(Exam.id.in_(exam_ids))).all() if exam_ids else []
    exam_map = {e.id: e for e in batch_exams}

    course_ids_from_batches = list({e.course_id for e in batch_exams})
    all_course_ids = list(set(course_ids_from_batches + [c.id for c in direct_courses]))

    if not all_course_ids:
        return []

    all_courses = db.scalars(select(Course).where(Course.id.in_(all_course_ids))).all()
    course_map = {c.id: c for c in all_courses}

    # Collect semester ids
    semester_ids = list({c.semester_id for c in all_courses if c.semester_id})
    semester_rows = db.scalars(select(Semester).where(Semester.id.in_(semester_ids))).all() if semester_ids else []
    semester_map = {s.id: s for s in semester_rows}

    # Group: semester_id -> semester data
    result_map: dict = {}

    for c in all_courses:
        sem_id = c.semester_id or "__no_semester__"
        sem = semester_map.get(sem_id) if sem_id != "__no_semester__" else None

        if sem_id not in result_map:
            result_map[sem_id] = {
                "semester_id": sem_id if sem_id != "__no_semester__" else None,
                "semester_name": sem.name if sem else "Chưa phân học kỳ",
                "semester_code": getattr(sem, "code", None),
                "semester_year": sem.year if sem else None,
                "semester_term": sem.term if sem else None,
                "semester_status": sem.status if sem else None,
                "start_date": sem.start_date if sem else None,
                "end_date": sem.end_date if sem else None,
                "courses": {},
            }

        result_map[sem_id]["courses"][c.id] = {
            "id": c.id,
            "code": c.code,
            "name": c.name,
            "description": c.description,
            "status": c.status,
            "exams": {},
        }

    # Attach exams and batches
    for batch in batches:
        exam = exam_map.get(batch.exam_id)
        if not exam:
            continue
        course = course_map.get(exam.course_id)
        if not course:
            continue

        sem_id = course.semester_id or "__no_semester__"
        if sem_id not in result_map or course.id not in result_map[sem_id]["courses"]:
            continue

        course_entry = result_map[sem_id]["courses"][course.id]
        if exam.id not in course_entry["exams"]:
            course_entry["exams"][exam.id] = {
                "id": exam.id,
                "name": exam.name,
                "status": exam.status,
                "time_limit": exam.time_limit,
                "question_count": exam.question_count,
                "batches": [],
            }

        slots = db.scalars(
            select(ScheduleSlot).where(ScheduleSlot.batch_id == batch.id).order_by(ScheduleSlot.slot_number)
        ).all()
        rooms = list(dict.fromkeys(s.room for s in slots if s.room))

        course_entry["exams"][exam.id]["batches"].append({
            "batch_id": batch.id,
            "name": batch.name,
            "date": batch.date,
            "start_time": batch.start_time,
            "end_time": batch.end_time,
            "total_assigned": batch.total_assigned,
            "status": batch.status,
            "slot_count": len(slots),
            "rooms": rooms,
        })

    # Also load exams for direct courses that might not have batches yet
    for sem_id, sem_data in result_map.items():
        for course_id, c_data in sem_data["courses"].items():
            if not c_data["exams"]:
                c_exams = db.scalars(select(Exam).where(Exam.course_id == course_id)).all()
                for e in c_exams:
                    c_data["exams"][e.id] = {
                        "id": e.id,
                        "name": e.name,
                        "status": e.status,
                        "time_limit": e.time_limit,
                        "question_count": e.question_count,
                        "batches": [],
                    }

    # Flatten to list
    output = []
    for sem_entry in result_map.values():
        courses_out = []
        for c_entry in sem_entry["courses"].values():
            exams_out = list(c_entry["exams"].values())
            courses_out.append({**c_entry, "exams": exams_out})
        output.append({**sem_entry, "courses": courses_out})

    def _sem_sort_key(s):
        status_order = {"ACTIVE": 0, "DRAFT": 1, "COMPLETED": 2}
        return (status_order.get(s.get("semester_status") or "", 9), -(s.get("semester_year") or 0))

    output.sort(key=_sem_sort_key)
    return output


@router.post("/courses", status_code=201)
def create_course(body: s.CourseIn, db: Session = Depends(get_db), user=Depends(editor)):
    row = Course(**body.model_dump(), owner_id=user.id)
    db.add(row)
    db.commit()
    return data(row, "code", "name", "description", "status")


@router.put("/courses/{course_id}")
def update_course(course_id: str, body: s.CourseIn, db: Session = Depends(get_db), user=Depends(editor)):
    row = course_access(db, course_id, user)
    for key, value in body.model_dump().items():
        setattr(row, key, value)
    db.commit()
    return data(row, "code", "name", "description", "status")


@router.delete("/courses/{course_id}")
def delete_course(
    course_id: str, body: s.CourseDeleteIn | None = None,
    db: Session = Depends(get_db), user=Depends(editor),
):
    if body is not None:
        if user.role not in {"SYSTEM_ADMIN", "ADMIN"}:
            fail(403, "FORBIDDEN", "Chỉ admin được xóa toàn bộ môn học và dữ liệu liên quan")
        return delete_course_tree(db, course_id, body.confirm_code, user)
    row = course_access(db, course_id, user)
    for model in (LearningOutcome, Topic, Document, Rubric, Exam, CourseEnrollment):
        if db.scalar(select(model.id).where(model.course_id == course_id).limit(1)):
            fail(
                409,
                "COURSE_IN_USE",
                "Môn học còn dữ liệu. Xóa dữ liệu liên quan trước hoặc lưu trữ môn học để giữ lịch sử.",
            )
    db.delete(row)
    db.commit()
    return {"ok": True}


@router.post("/courses/{course_id}/archive")
def archive_course(course_id: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = course_access(db, course_id, user)
    row.status = "ARCHIVED"
    db.commit()
    return {"status": "ARCHIVED"}


@router.post("/courses/{course_id}/restore")
def restore_course(course_id: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = course_access(db, course_id, user)
    row.status = "ACTIVE"
    db.commit()
    return {"status": "ACTIVE"}


@router.get("/courses/{course_id}/workspace")
def workspace(course_id: str, db: Session = Depends(get_db), user=Depends(staff)):
    course_access(db, course_id, user)

    def rows(model, *fields):
        return [
            data(r, *fields)
            for r in db.scalars(select(model).where(model.course_id == course_id).order_by(model.created_at))
        ]

    return {
        "outcomes": rows(LearningOutcome, "code", "description", "weight"),
        "topics": [
            topic_data(db, t)
            for t in db.scalars(select(Topic).where(Topic.course_id == course_id).order_by(Topic.created_at))
        ],
        "documents": [],
        "chapters": [],
        "rubrics": rows(Rubric, "name", "version", "criteria"),
        "exams": [
            data(exam, "name", "status", "blueprint", "time_limit", "question_count", "rubric_id", "max_attempts")
            | {"questions": [
                {"text": q["text"], "english_terms": q.get("english_terms", [])}
                for q in (exam.snapshot or {}).get("questions", [])
            ]}
            for exam in db.scalars(select(Exam).where(Exam.course_id == course_id).order_by(Exam.created_at))
        ],
        "items": [
            {
                "id": it.id,
                "topic_id": it.topic_id,
                "learning_outcome_id": it.learning_outcome_id,
                "difficulty": it.difficulty,
                "prompt": it.prompt,
                "expected_points": it.expected_points or [],
                "key_terms": it.key_terms or [],
                "status": it.status,
                "usage_count": it.usage_count,
            }
            for it in db.scalars(
                select(QuestionItem)
                .where(QuestionItem.course_id == course_id, QuestionItem.deleted_at.is_(None))
                .order_by(QuestionItem.created_at)
            )
        ],
    }


@router.post("/courses/{course_id}/outcomes", status_code=201)
def create_outcome(course_id: str, body: s.LOIn, db: Session = Depends(get_db), user=Depends(editor)):
    course_access(db, course_id, user)
    row = LearningOutcome(course_id=course_id, **body.model_dump())
    db.add(row)
    db.commit()
    return data(row, "code", "description", "weight")


@router.put("/outcomes/{key}")
def update_outcome(key: str, body: s.LOIn, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, LearningOutcome, key)
    course_access(db, row.course_id, user)
    for field, value in body.model_dump().items():
        setattr(row, field, value)
    db.commit()
    return data(row, "code", "description", "weight")


@router.delete("/outcomes/{key}")
def delete_outcome(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, LearningOutcome, key)
    course_access(db, row.course_id, user)
    db.delete(row)
    db.commit()
    return {"ok": True}


@router.post("/courses/{course_id}/topics", status_code=201)
def create_topic(course_id: str, body: s.TopicIn, db: Session = Depends(get_db), user=Depends(editor)):
    course_access(db, course_id, user)
    row = Topic(course_id=course_id, learning_outcome_id=body.learning_outcome_ids[0], name=body.name)
    db.add(row)
    set_mappings(db, row, body)
    db.commit()
    return topic_data(db, row)


@router.put("/topics/{key}")
def update_topic(key: str, body: s.TopicIn, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, Topic, key, lock=True)
    course_access(db, row.course_id, user)
    set_mappings(db, row, body)
    db.commit()
    return topic_data(db, row)


@router.delete("/topics/{key}")
def delete_topic(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, Topic, key)
    course_access(db, row.course_id, user)
    if any(
        any(b["topic_id"] == key for b in e.blueprint)
        for e in db.scalars(select(Exam).where(Exam.course_id == row.course_id))
    ):
        fail(409, "IN_USE", "Chủ đề đang được sử dụng trong bài thi")
    db.delete(row)
    db.commit()
    return {"ok": True}


# === Item Bank (Ngân hàng câu hỏi) Endpoints ===

@router.get("/courses/{course_id}/items")
def list_question_items(
    course_id: str,
    topic_id: str | None = None,
    difficulty: str | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    course_access(db, course_id, user)
    stmt = select(QuestionItem).where(QuestionItem.course_id == course_id, QuestionItem.deleted_at.is_(None))
    if topic_id:
        stmt = stmt.where(QuestionItem.topic_id == topic_id)
    if difficulty:
        stmt = stmt.where(QuestionItem.difficulty == difficulty)
    if status:
        stmt = stmt.where(QuestionItem.status == status)
    items = db.scalars(stmt.order_by(QuestionItem.created_at.desc())).all()
    return [
        {
            "id": it.id,
            "topic_id": it.topic_id,
            "learning_outcome_id": it.learning_outcome_id,
            "difficulty": it.difficulty,
            "prompt": it.prompt,
            "expected_points": it.expected_points or [],
            "key_terms": it.key_terms or [],
            "status": it.status,
            "usage_count": it.usage_count,
            "created_at": it.created_at,
        }
        for it in items
    ]


@router.post("/courses/{course_id}/items", status_code=201)
def create_question_item(
    course_id: str,
    body: s.QuestionBankItemIn,
    db: Session = Depends(get_db),
    user=Depends(editor),
):
    course_access(db, course_id, user)
    topic = by_id(db, Topic, body.topic_id)
    if topic.course_id != course_id:
        fail(400, "INVALID_TOPIC", "Chủ đề không thuộc môn học này")
    item = QuestionItem(
        course_id=course_id,
        topic_id=body.topic_id,
        learning_outcome_id=body.learning_outcome_id or topic.learning_outcome_id,
        difficulty=body.difficulty,
        prompt=body.prompt.strip(),
        expected_points=[p.strip() for p in body.expected_points if p.strip()],
        key_terms=[t.strip() for t in body.key_terms if t.strip()],
        status=body.status,
        author_id=user.id,
    )
    db.add(item)
    db.commit()
    return {
        "id": item.id,
        "topic_id": item.topic_id,
        "learning_outcome_id": item.learning_outcome_id,
        "difficulty": item.difficulty,
        "prompt": item.prompt,
        "expected_points": item.expected_points,
        "key_terms": item.key_terms,
        "status": item.status,
    }


@router.post("/courses/{course_id}/items/import", status_code=201)
def import_question_items(
    course_id: str,
    body: s.QuestionBankItemImportIn,
    db: Session = Depends(get_db),
    user=Depends(editor),
):
    course_access(db, course_id, user)
    created = []
    for q in body.items:
        topic = db.get(Topic, q.topic_id)
        if not topic or topic.course_id != course_id:
            continue
        item = QuestionItem(
            course_id=course_id,
            topic_id=q.topic_id,
            learning_outcome_id=q.learning_outcome_id or topic.learning_outcome_id,
            difficulty=q.difficulty,
            prompt=q.prompt.strip(),
            expected_points=[p.strip() for p in q.expected_points if p.strip()],
            key_terms=[t.strip() for t in q.key_terms if t.strip()],
            status=q.status,
            author_id=user.id,
        )
        db.add(item)
        created.append(item)
    db.commit()
    return {"imported_count": len(created)}


@router.get("/courses/{course_id}/items/export-template")
def export_question_items_template(
    course_id: str,
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse

    course = course_access(db, course_id, user)
    topics = db.scalars(select(Topic).where(Topic.course_id == course_id)).all()
    los = db.scalars(select(LearningOutcome).where(LearningOutcome.course_id == course_id)).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ngan_Hang_Cau_Hoi"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "Tên hoặc Mã Chủ đề (*)",
        "Mã Chuẩn đầu ra (LO)",
        "Độ khó (EASY / MEDIUM / HARD)",
        "Nội dung câu hỏi thi vấn đáp (*)",
        "Điểm mấu chốt (cách nhau bởi dấu ;)",
        "Từ khóa chuyên môn (cách nhau bởi dấu ,)",
        "Trạng thái (APPROVED / DRAFT)",
    ]
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    sample_topic = topics[0].name if topics else "Cấu trúc dữ liệu nâng cao"
    sample_lo = los[0].code if los else "LO1"
    sample_rows = [
        [
            sample_topic,
            sample_lo,
            "MEDIUM",
            "Trình bày điều kiện cân bằng của cây AVL và giải thích cơ chế quay đơn khi bị mất cân bằng.",
            "Cây nhị phân tìm kiếm; Hệ số cân bằng Balance Factor trong khoảng -1, 0, 1; Thực hiện phép quay trái hoặc quay phải khi mất cân bằng tại nút cha",
            "AVL Tree, Balance Factor, Single Rotation, Left-Right Rotation",
            "APPROVED",
        ],
        [
            sample_topic,
            sample_lo,
            "HARD",
            "So sánh độ phức tạp thời gian và không gian giữa thuật toán Dijkstra và Bellman-Ford trong tìm đường đi ngắn nhất.",
            "Dijkstra dùng hàng đợi ưu tiên O((V+E)logV), chỉ áp dụng cho trọng số không âm; Bellman-Ford O(V*E), phát hiện được chu trình âm",
            "Dijkstra, Bellman-Ford, Shortest Path, Negative Cycle, Time Complexity",
            "APPROVED",
        ]
    ]
    for row in sample_rows:
        ws.append(row)

    col_widths = [28, 20, 22, 55, 45, 38, 18]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: Danh mục tham chiếu
    ws2 = wb.create_sheet(title="Danh_Muc_Tham_Chieu")
    ws2.append(["DANH MỤC CHỦ ĐỀ HIỆN CÓ CỦA MÔN", "", "DANH MỤC CHUẨN ĐẦU RA (LO)"])
    ws2.append(["Tên chủ đề", "ID chủ đề", "Mã LO", "Mô tả LO"])
    ws2.row_dimensions[1].height = 24
    ws2.row_dimensions[2].height = 20

    ref_header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    for c in range(1, 5):
        cell = ws2.cell(row=2, column=c)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = ref_header_fill

    max_len = max(len(topics), len(los), 1)
    for idx in range(max_len):
        t_name = topics[idx].name if idx < len(topics) else ""
        t_id = topics[idx].id if idx < len(topics) else ""
        l_code = los[idx].code if idx < len(los) else ""
        l_desc = los[idx].description if idx < len(los) else ""
        ws2.append([t_name, t_id, l_code, l_desc])

    ws2.column_dimensions['A'].width = 30
    ws2.column_dimensions['B'].width = 38
    ws2.column_dimensions['C'].width = 15
    ws2.column_dimensions['D'].width = 45

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Template_Ngan_Hang_Cau_Hoi_{course.code}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/courses/{course_id}/items/import-excel", status_code=201)
async def import_question_items_excel(
    course_id: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(editor),
):
    import io
    import openpyxl

    course_access(db, course_id, user)
    content = await file.read()
    if not content:
        fail(400, "EMPTY_FILE", "File tải lên không có dữ liệu")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        fail(400, "INVALID_EXCEL", f"Không thể đọc file Excel: {str(e)}")

    ws = wb.active
    topics = db.scalars(select(Topic).where(Topic.course_id == course_id)).all()
    topic_map = {}
    for t in topics:
        topic_map[t.id.lower()] = t.id
        topic_map[t.name.lower().strip()] = t.id

    los = db.scalars(select(LearningOutcome).where(LearningOutcome.course_id == course_id)).all()
    lo_map = {}
    for l in los:
        lo_map[l.id.lower()] = l.id
        lo_map[l.code.lower().strip()] = l.id

    imported_items = []
    errors = []

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        fail(400, "EMPTY_DATA", "File không có dữ liệu câu hỏi (tối thiểu cần 1 dòng sau tiêu đề)")

    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or not any(row):
            continue

        topic_raw = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
        lo_raw = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        diff_raw = str(row[2]).strip().upper() if len(row) > 2 and row[2] is not None else "MEDIUM"
        prompt_raw = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        expected_raw = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
        terms_raw = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
        status_raw = str(row[6]).strip().upper() if len(row) > 6 and row[6] is not None else "APPROVED"

        if not prompt_raw:
            errors.append(f"Dòng {row_idx}: Nội dung câu hỏi không được để trống")
            continue

        topic_id = topic_map.get(topic_raw.lower())
        if not topic_id:
            if topics:
                topic_id = topics[0].id
            else:
                errors.append(f"Dòng {row_idx}: Không tìm thấy chủ đề '{topic_raw}' trong môn học")
                continue

        lo_id = lo_map.get(lo_raw.lower()) if lo_raw else None
        if not lo_id:
            matched_topic = next((t for t in topics if t.id == topic_id), None)
            lo_id = matched_topic.learning_outcome_id if matched_topic else None

        difficulty = diff_raw if diff_raw in {"EASY", "MEDIUM", "HARD"} else "MEDIUM"
        status = "DRAFT" if status_raw == "DRAFT" else "APPROVED"

        expected_points = [p.strip() for p in expected_raw.split(";") if p.strip()] if expected_raw else []
        if not expected_points and "\n" in expected_raw:
            expected_points = [p.strip() for p in expected_raw.split("\n") if p.strip()]

        key_terms = []
        if terms_raw:
            sep = ";" if ";" in terms_raw else ","
            key_terms = [t.strip() for t in terms_raw.split(sep) if t.strip()]

        item = QuestionItem(
            course_id=course_id,
            topic_id=topic_id,
            learning_outcome_id=lo_id,
            difficulty=difficulty,
            prompt=prompt_raw,
            expected_points=expected_points,
            key_terms=key_terms,
            status=status,
            author_id=user.id,
        )
        db.add(item)
        imported_items.append(item)

    if imported_items:
        db.commit()

    return {
        "imported_count": len(imported_items),
        "errors": errors,
        "total_rows_processed": len(rows) - 1,
    }


@router.delete("/items/{key}")
def delete_question_item(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    item = by_id(db, QuestionItem, key)
    course_access(db, item.course_id, user)
    item.deleted_at = time.time()
    db.commit()
    return {"ok": True}


@router.post("/courses/{course_id}/documents", status_code=201)
async def upload_document(
    course_id: str,
    topic_id: str | None = Form(default=None),
    kind: str = Form(default="SUPPLEMENT"),
    file: UploadFile = File(),
    db: Session = Depends(get_db),
    user=Depends(editor),
):
    course_access(db, course_id, user)
    if kind not in {"TEXTBOOK", "SUPPLEMENT"}:
        fail(422, "INVALID_KIND", "Loại tài liệu không hợp lệ")
    if topic_id and by_id(db, Topic, topic_id).course_id != course_id:
        fail(422, "CROSS_COURSE", "Chủ đề không thuộc môn học")
    filename = Path(file.filename or "").name[:250]
    ext = filename.rsplit(".", 1)[-1].lower()
    if kind == "TEXTBOOK":
        if ext != "pdf" or topic_id:
            fail(422, "TEXTBOOK_PDF", "Giáo trình là PDF dùng chung cho môn, không chọn chủ đề")
        if db.scalar(select(Document.id).where(Document.course_id == course_id, Document.kind == "TEXTBOOK")):
            fail(409, "TEXTBOOK_EXISTS", "Môn học đã có giáo trình PDF")
    if ext not in {"pdf", "pptx", "docx", "txt"}:
        fail(422, "INVALID_FORMAT", "Chỉ nhận PDF, PPTX, DOCX, TXT")
    limit = (settings().max_textbook_mb if kind == "TEXTBOOK" else settings().max_document_mb) * 1024 * 1024
    content = await file.read(limit + 1)
    if not content or len(content) > limit:
        fail(413, "DOCUMENT_SIZE", "Tài liệu rỗng hoặc quá lớn")
    key = f"documents/{uid()}/{filename}"
    storage.put(key, content)
    row = Document(
        course_id=course_id,
        topic_id=topic_id,
        kind=kind,
        filename=filename,
        storage_key=key,
        embedding_model=ai.embedding_name(),
    )
    db.add(row)
    db.flush()
    if topic_id:
        db.add(TopicDocument(topic_id=topic_id, document_id=row.id))
    db.add(Audit(user_id=user.id, event="DOCUMENT_UPLOADED", details={"course_id": course_id}))
    db.commit()
    return data(row, "filename", "status")


@router.get("/documents/{key}/content")
def document_content(key: str, db: Session = Depends(get_db), user=Depends(staff)):
    document = by_id(db, Document, key)
    course_access(db, document.course_id, user)
    return Response(
        storage.get(document.storage_key),
        media_type="application/octet-stream",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(document.filename, safe="")},
    )


@router.put("/documents/{key}/file")
async def replace_failed_textbook(
    key: str, file: UploadFile = File(), db: Session = Depends(get_db), user=Depends(editor)
):
    row = by_id(db, Document, key, lock=True)
    course_access(db, row.course_id, user)
    if row.kind != "TEXTBOOK" or row.status != "FAILED":
        fail(
            409,
            "INVALID_STATE",
            "Chỉ thay PDF giáo trình xử lý lỗi; giáo trình đã dùng phải được giữ để đối chiếu",
        )
    name = Path(file.filename or "").name[:250]
    limit = settings().max_textbook_mb * 1024 * 1024
    content = await file.read(limit + 1)
    if not name.lower().endswith(".pdf") or not content.startswith(b"%PDF-") or len(content) > limit:
        fail(422, "INVALID_PDF", "Cần file PDF hợp lệ trong giới hạn dung lượng giáo trình")
    storage_key = f"documents/{uid()}/{name}"
    storage.put(storage_key, content)
    row.storage_key, row.filename = storage_key, name
    row.status, row.error, row.version = "PENDING", None, row.version + 1
    db.add(
        Audit(
            user_id=user.id,
            event="FAILED_TEXTBOOK_REPLACED",
            details={"document_id": key, "version": row.version},
        )
    )
    db.commit()
    return data(row, "filename", "status", "version")


@router.post("/documents/{key}/chapters", status_code=201)
def add_section(key: str, body: s.BookSectionIn, db: Session = Depends(get_db), user=Depends(editor)):
    doc = by_id(db, Document, key)
    course_access(db, doc.course_id, user)
    if doc.kind != "TEXTBOOK" or doc.status != "READY" or body.end_page > (doc.page_count or 0):
        fail(422, "INVALID_PAGES", "Chọn số trang PDF hợp lệ của giáo trình đã xử lý")
    row = BookSection(course_id=doc.course_id, document_id=doc.id, **body.model_dump(), source="MANUAL")
    db.add(row)
    db.commit()
    return data(row, "title", "level", "start_page", "end_page", "source")


@router.put("/chapters/{key}")
def update_section(key: str, body: s.BookSectionIn, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, BookSection, key, lock=True)
    course_access(db, row.course_id, user)
    if body.end_page > (by_id(db, Document, row.document_id).page_count or 0):
        fail(422, "INVALID_PAGES", "Trang vượt số trang PDF")
    for field, value in body.model_dump().items():
        setattr(row, field, value)
    row.source = "MANUAL"
    db.commit()
    return data(row, "title", "level", "start_page", "end_page", "source")


@router.delete("/chapters/{key}")
def delete_section(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, BookSection, key)
    course_access(db, row.course_id, user)
    db.delete(row)
    db.commit()
    return {"ok": True}


@router.post("/documents/{key}/retry")
def retry_document(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, Document, key, lock=True)
    course_access(db, row.course_id, user)
    if row.status != "FAILED":
        fail(409, "INVALID_STATE", "Chỉ xử lý lại tài liệu lỗi")
    row.status, row.error, row.embedding_model = "PENDING", None, ai.embedding_name()
    db.commit()
    return {"status": row.status}


@router.get("/courses/{course_id}/rag")
def rag_debug(course_id: str, topic_id: str, q: str, db: Session = Depends(get_db), user=Depends(staff)):
    course_access(db, course_id, user)
    if not q.strip() or len(q) > 3000:
        fail(422, "INVALID_QUERY", "Truy vấn phải từ 1–3000 ký tự")
    return ai.retrieve(db, course_id, topic_id, q)


@router.post("/courses/{course_id}/rubrics", status_code=201)
def create_rubric(course_id: str, body: s.RubricIn, db: Session = Depends(get_db), user=Depends(editor)):
    course_access(db, course_id, user)
    row = Rubric(course_id=course_id, **body.model_dump())
    db.add(row)
    db.commit()
    return data(row, "name", "criteria", "version")


@router.put("/rubrics/{key}")
def update_rubric(key: str, body: s.RubricIn, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, Rubric, key, lock=True)
    course_access(db, row.course_id, user)
    row.name, row.criteria, row.version = body.name, body.model_dump()["criteria"], row.version + 1
    db.commit()
    return data(row, "name", "criteria", "version")


@router.delete("/rubrics/{key}")
def delete_rubric(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, Rubric, key, lock=True)
    course_access(db, row.course_id, user)
    if db.scalar(select(Exam.id).where(Exam.rubric_id == key).limit(1)):
        fail(
            409,
            "RUBRIC_IN_USE",
            "Rubric đang được đề thi sử dụng. Đổi rubric hoặc xóa đề nháp liên quan trước; đề đã công bố giữ nguyên lịch sử.",
        )
    db.delete(row)
    db.commit()
    return {"ok": True}


def validate_exam(db, body, user):
    course = course_access(db, body.course_id, user)
    if course.status == "ARCHIVED":
        fail(409, "ARCHIVED", "Môn học đã lưu trữ")
    if by_id(db, Rubric, body.rubric_id).course_id != body.course_id:
        fail(422, "CROSS_COURSE", "Rubric không thuộc môn học")
    for row in body.blueprint:
        if by_id(db, Topic, row.topic_id).course_id != body.course_id:
            fail(422, "CROSS_COURSE", "Blueprint chứa chủ đề không thuộc môn học")


@router.post("/exams", status_code=201)
def create_exam(body: s.ExamIn, db: Session = Depends(get_db), user=Depends(editor)):
    validate_exam(db, body, user)
    if user.role not in {"SYSTEM_ADMIN", "ADMIN", "EXAMINER", "REVIEWER"} and body.max_attempts != 1:
        fail(403, "FORBIDDEN", "Chỉ admin được cấu hình số lượt làm bài")
    row = Exam(**body.model_dump())
    db.add(row)
    db.commit()
    return data(row, "name", "status")


@router.put("/exams/{key}")
def update_exam(key: str, body: s.ExamIn, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, Exam, key, lock=True)
    course_access(db, row.course_id, user)
    if row.status != "DRAFT":
        fail(409, "PUBLISHED", "Không sửa đề đã công bố")
    if body.course_id != row.course_id:
        fail(422, "CROSS_COURSE", "Không chuyển đề thi sang môn học khác")
    validate_exam(db, body, user)
    if user.role not in {"SYSTEM_ADMIN", "ADMIN", "EXAMINER", "REVIEWER"} and body.max_attempts != row.max_attempts:
        fail(403, "FORBIDDEN", "Chỉ admin được cấu hình số lượt làm bài")
    for field, value in body.model_dump().items():
        setattr(row, field, value)
    db.commit()
    return data(row, "name", "status", "max_attempts")


@router.delete("/exams/{key}")
def delete_exam(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    row = by_id(db, Exam, key, lock=True)
    course_access(db, row.course_id, user)
    if row.status != "DRAFT":
        fail(409, "PUBLISHED", "Không xóa đề đã công bố")
    db.delete(row)
    db.commit()
    return {"ok": True}


@router.post("/exams/{key}/copy", status_code=201)
def copy_exam(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    source = by_id(db, Exam, key)
    body = s.ExamIn(
        course_id=source.course_id,
        rubric_id=source.rubric_id,
        name=f"{source.name[:189]} (bản sao)",
        time_limit=source.time_limit,
        blueprint=deepcopy(source.blueprint),
        max_attempts=source.max_attempts,
    )
    validate_exam(db, body, user)
    row = Exam(**body.model_dump())
    db.add(row)
    db.commit()
    return data(row, "name", "status", "blueprint", "time_limit", "rubric_id", "max_attempts")


@router.post("/exams/{key}/publish")
def publish(key: str, db: Session = Depends(get_db), user=Depends(editor)):
    exam = by_id(db, Exam, key, lock=True)
    course_access(db, exam.course_id, user)
    if exam.status == "PUBLISHED":
        return {"status": exam.status}

    rubric = db.get(Rubric, exam.rubric_id) if exam.rubric_id else None
    if not rubric:
        rubric = db.scalar(
            select(Rubric).where(Rubric.course_id == exam.course_id).order_by(Rubric.created_at.desc())
        )

    questions = []
    blueprint = exam.blueprint or []

    # Lắp ráp đề thi từ Ngân hàng câu hỏi (Item Bank) dựa theo Ma trận chuẩn (Blueprint)
    if not blueprint:
        # Nếu chưa có blueprint chi tiết, bốc từ tất cả câu hỏi APPROVED của môn
        available_items = db.scalars(
            select(QuestionItem).where(
                QuestionItem.course_id == exam.course_id,
                QuestionItem.status == "APPROVED",
                QuestionItem.deleted_at.is_(None),
            ).order_by(QuestionItem.usage_count.asc())
        ).all()
        target_count = exam.question_count or 3
        if len(available_items) < target_count:
            fail(
                400,
                "INSUFFICIENT_ITEMS",
                f"Ngân hàng câu hỏi của môn chỉ có {len(available_items)} câu khả dụng, cần tối thiểu {target_count} câu để tạo đề thi. Vui lòng nạp thêm câu hỏi.",
            )
        selected_items = random.sample(available_items, target_count)
        for item in selected_items:
            questions.append({
                "item_id": item.id,
                "text": item.prompt,
                "expected_points": item.expected_points or [],
                "key_terms": item.key_terms or [],
                "english_terms": [{"term": t, "meaning": ""} for t in (item.key_terms or [])],
                "topic_id": item.topic_id,
                "learning_outcome_id": item.learning_outcome_id,
                "difficulty": item.difficulty,
            })
            item.usage_count += 1
    else:
        for row in blueprint:
            topic_id = row.get("topic_id")
            difficulty = row.get("difficulty")
            count = int(row.get("count", 1))

            query = select(QuestionItem).where(
                QuestionItem.course_id == exam.course_id,
                QuestionItem.status == "APPROVED",
                QuestionItem.deleted_at.is_(None),
            )
            if topic_id:
                query = query.where(QuestionItem.topic_id == topic_id)
            if difficulty:
                query = query.where(QuestionItem.difficulty == difficulty)

            matched_items = db.scalars(query.order_by(QuestionItem.usage_count.asc())).all()

            # Nếu không đủ câu hỏi đúng mức độ khó yêu cầu, nới lỏng sang các câu hỏi khác trong cùng topic
            if len(matched_items) < count and topic_id:
                matched_items = db.scalars(
                    select(QuestionItem).where(
                        QuestionItem.course_id == exam.course_id,
                        QuestionItem.topic_id == topic_id,
                        QuestionItem.status == "APPROVED",
                        QuestionItem.deleted_at.is_(None),
                    ).order_by(QuestionItem.usage_count.asc())
                ).all()

            if len(matched_items) < count:
                topic = db.get(Topic, topic_id) if topic_id else None
                topic_name = topic.name if topic else "Đã chọn"
                fail(
                    400,
                    "INSUFFICIENT_ITEMS",
                    f"Chủ đề '{topic_name}' chỉ có {len(matched_items)} câu hỏi khả dụng trong Ngân hàng, không đủ {count} câu theo ma trận đề thi. Vui lòng bổ sung thêm câu hỏi vào Ngân hàng.",
                )

            selected_items = random.sample(matched_items, count)
            for item in selected_items:
                questions.append({
                    "item_id": item.id,
                    "text": item.prompt,
                    "expected_points": item.expected_points or [],
                    "key_terms": item.key_terms or [],
                    "english_terms": [{"term": t, "meaning": ""} for t in (item.key_terms or [])],
                    "topic_id": item.topic_id,
                    "learning_outcome_id": item.learning_outcome_id,
                    "difficulty": item.difficulty,
                })
                item.usage_count += 1

    criteria = (
        rubric.criteria
        if rubric
        else [
            {"name": "Độ chính xác kiến thức", "max_score": 5.0, "weight": 0.5},
            {"name": "Lập luận và bản chất", "max_score": 3.0, "weight": 0.3},
            {"name": "Thuật ngữ chuyên ngành", "max_score": 2.0, "weight": 0.2},
        ]
    )

    exam.snapshot = {
        "exam_version": 3,
        "assembly_mode": "ITEM_BANK_BLUEPRINT",
        "rubric_id": rubric.id if rubric else None,
        "rubric_version": rubric.version if rubric else 1,
        "criteria": criteria,
        "questions": questions,
        "ai_provider": settings().ai_provider,
        "llm_model": settings().llm_model,
        "prompt_version": ai.PROMPT_VERSION,
        "knowledge_version": "item-bank-v1",
        "published_at": time.time(),
    }
    exam.status = "PUBLISHED"
    db.add(Audit(user_id=user.id, event="EXAM_PUBLISHED", details={"exam_id": exam.id, "questions_count": len(questions)}))
    db.commit()
    return {"status": exam.status, "question_count": len(questions)}


@router.post("/exams/{key}/assign")
def assign(key: str, body: s.AssignIn, db: Session = Depends(get_db), user=Depends(editor)):
    exam = by_id(db, Exam, key, lock=True)
    course_access(db, exam.course_id, user)
    if exam.status != "PUBLISHED":
        fail(409, "NOT_PUBLISHED", "Công bố đề trước khi giao bài")
    for student_id in set(body.student_ids):
        target = by_id(db, User, student_id)
        if target.role != "STUDENT" or target.status != "ACTIVE":
            fail(422, "NOT_STUDENT", "Chỉ giao bài cho sinh viên đang hoạt động")
        if not db.scalar(
            select(Assignment).where(Assignment.exam_id == key, Assignment.student_id == student_id)
        ):
            db.add(Assignment(exam_id=key, student_id=student_id))
    db.commit()
    return {"ok": True}


STAGE_LABELS = {
    "READY": "Chưa làm",
    "SUBMITTED": "Đang xếp hàng chờ Worker...",
    "TRANSCRIBING": "Đang phiên âm Whisper STT...",
    "ANALYZING": "Đang phân tích độ trôi chảy & âm học...",
    "GRADING": "AI đang chấm điểm theo Rubric...",
    "GRADED": "Đã chấm xong câu hỏi",
}

STAGE_WEIGHTS = {
    "READY": 0,
    "SUBMITTED": 10,
    "TRANSCRIBING": 35,
    "ANALYZING": 60,
    "GRADING": 85,
    "GRADED": 100,
}


def calculate_session_progress(db: Session, session: ExamSession):
    if session.status in {"COMPLETED", "APPROVED"}:
        return None
    attempts = db.scalars(
        select(Attempt).where(Attempt.session_id == session.id).order_by(Attempt.sequence)
    ).all()
    if not attempts:
        return None

    total = len(attempts)
    completed = sum(1 for a in attempts if a.assessment is not None or a.status == "GRADED")
    active_attempt = next((a for a in attempts if a.assessment is None), None)

    if active_attempt:
        current_stage = active_attempt.status or "SUBMITTED"
        current_seq = active_attempt.sequence
    else:
        current_stage = "GRADED"
        current_seq = total

    stage_label = STAGE_LABELS.get(current_stage, "Đang xử lý AI...")

    total_points = sum(
        100 if a.assessment is not None or a.status == "GRADED" else STAGE_WEIGHTS.get(a.status, 10)
        for a in attempts
    )
    percent = min(99, max(5, int(total_points / total))) if session.status == "SUBMITTED" else 0

    return {
        "total_questions": total,
        "completed_questions": completed,
        "current_sequence": current_seq,
        "current_stage": current_stage,
        "stage_label": stage_label,
        "percent": percent,
    }


@router.get("/results")
def results(db: Session = Depends(get_db), user=Depends(staff)):
    query = (
        select(ExamSession, Exam, User)
        .join(Exam, Exam.id == ExamSession.exam_id)
        .join(User, User.id == ExamSession.student_id)
        .where(ExamSession.deleted_at.is_(None))
    )
    if user.role == "TEACHER":
        query = query.join(Course, Course.id == Exam.course_id).where(
            or_(
                Course.owner_id == user.id,
                Course.teacher_id == user.id,
                Exam.id.in_(
                    select(ExamBatch.exam_id).where(ExamBatch.assigned_teacher_id == user.id)
                ),
                Course.id.in_(
                    select(Section.course_id).where(Section.teacher_id == user.id)
                ),
                Course.code == "ORAL-PRACTICE",
            )
        )
    session_rows = []
    for session, exam, student in db.execute(query.order_by(ExamSession.created_at.desc())):
        row = (
            history_row(session)
            | {"exam_name": exam.name, "student_name": student.name,
               "student_id": student.id, "exam_id": exam.id, **allowance(db, exam, student.id)}
        )
        if session.status == "SUBMITTED":
            row["progress"] = calculate_session_progress(db, session)
        session_rows.append(row)
    return session_rows


@router.get("/results/{key}")
def review(key: str, db: Session = Depends(get_db), user=Depends(staff)):
    session = by_id(db, ExamSession, key)
    if session.deleted_at is not None:
        fail(404, "NOT_FOUND", "Lần thi đã bị xóa")
    exam = by_id(db, Exam, session.exam_id)
    course_access(db, exam.course_id, user)
    attempts = db.scalars(select(Attempt).where(Attempt.session_id == key).order_by(Attempt.sequence)).all()
    session_progress = calculate_session_progress(db, session) if session.status == "SUBMITTED" else None
    return history_row(session) | {
        "progress": session_progress,
        "exam_id": exam.id, "student_id": session.student_id,
        **allowance(db, exam, session.student_id),
        "history": [history_row(row) for row in sessions_for(db, exam.id, session.student_id)],
        "exam_name": exam.name,
        "snapshot": exam.snapshot,
        "student_name": by_id(db, User, session.student_id).name,
        "attempts": [
            data(a, "sequence", "question", "transcript", "stt_confidence", "assessment", "status")
            | {
                "stage": a.status,
                "stage_label": STAGE_LABELS.get(a.status, a.status),
                "assessment": assessment_view(a.assessment, exam),
                "grading_targets": grading_targets(db, exam, a),
                "evidence": [
                    data(e, "id", "kind", "status", "sha256", "size")
                    for e in db.scalars(
                        select(Upload).where(Upload.attempt_id == a.id, Upload.status == "COMPLETED")
                    )
                ],
                "reviews": [
                    data(
                        j,
                        "status",
                        "reason",
                        "policy",
                        "original",
                        "result",
                        "error",
                        "created_at",
                        "completed_at",
                        "requested_by",
                    )
                    for j in db.scalars(
                        select(ReviewJob)
                        .where(ReviewJob.attempt_id == a.id)
                        .order_by(ReviewJob.created_at.desc())
                    )
                ],
            }
            for a in attempts
        ],
    }


def grading_targets(db, exam, attempt):
    targets = []
    for candidate in db.scalars(select(Exam).where(Exam.course_id == exam.course_id, Exam.status == "PUBLISHED")):
        try:
            review_question(exam, candidate, attempt)
        except GradingError:
            continue
        targets.append({"id": candidate.id, "name": candidate.name,
                        "model": candidate.snapshot["llm_model"]})
    return targets


@router.post("/attempts/{key}/grade-review", status_code=202)
def grade_review(key: str, body: s.GradeReviewIn, db: Session = Depends(get_db), user=Depends(admin)):
    attempt = by_id(db, Attempt, key, lock=True)
    session = by_id(db, ExamSession, attempt.session_id, lock=True)
    if session.deleted_at is not None:
        fail(404, "NOT_FOUND", "Lần thi đã bị xóa")
    exam = by_id(db, Exam, session.exam_id)
    course_access(db, exam.course_id, user)
    if session.status not in {"SUBMITTED", "REVIEW_REQUIRED", "COMPLETED"} or attempt.status != "GRADED":
        fail(409, "NOT_FINISHED", "Chờ nộp bài và hoàn tất xử lý lần đầu")
    if not attempt.transcript or not attempt.finished_at:
        fail(409, "NO_TRANSCRIPT", "Chưa có transcript đã nộp để chấm lại")
    target = by_id(db, Exam, body.target_exam_id)
    try:
        review_question(exam, target, attempt)
    except GradingError as exc:
        fail(409, exc.code, str(exc))
    pending = db.scalar(select(ReviewJob).where(ReviewJob.attempt_id == key, ReviewJob.status == "PENDING"))
    if pending:
        if pending.policy == {"provider": "grading", "target_exam_id": target.id}:
            return data(pending, "status")
        fail(409, "REVIEW_PENDING", "Đang có yêu cầu xử lý khác; hãy chờ hoàn tất")
    job = ReviewJob(attempt_id=key, requested_by=user.id, reason=body.reason,
                    policy={"provider": "grading", "target_exam_id": target.id},
                    original={"transcript": attempt.transcript, "stt_confidence": attempt.stt_confidence,
                              "assessment": attempt.assessment, "source_exam_id": exam.id})
    db.add(job)
    session.status, session.final_score = "REVIEW_REQUIRED", None
    db.flush()
    db.add(Audit(user_id=user.id, event="GRADING_REVIEW_REQUESTED",
                 details={"job_id": job.id, "attempt_id": key, "target_exam_id": target.id, "reason": body.reason}))
    db.commit()
    return data(job, "status")


@router.post("/attempts/{key}/google-review", status_code=202)
def google_review(key: str, body: s.ReviewIn, db: Session = Depends(get_db), user=Depends(admin)):
    return request_transcription_review(key, body, db, user, "google")


@router.post("/attempts/{key}/gemini-review", status_code=202)
def gemini_review(key: str, body: s.ReviewIn, db: Session = Depends(get_db), user=Depends(admin)):
    return request_transcription_review(key, body, db, user, "gemini")


def request_transcription_review(key, body, db, user, provider):
    attempt = by_id(db, Attempt, key, lock=True)
    session = by_id(db, ExamSession, attempt.session_id, lock=True)
    exam = by_id(db, Exam, session.exam_id)
    course_access(db, exam.course_id, user)
    try:
        check_config(exam)
    except GradingError as exc:
        fail(409, exc.code, str(exc))
    if session.status not in {"SUBMITTED", "REVIEW_REQUIRED", "COMPLETED"} or attempt.status != "GRADED":
        fail(409, "NOT_FINISHED", "Chờ sinh viên nộp bài và hoàn tất chấm lần đầu")
    if provider == "gemini" and not settings().gemini_api_key:
        fail(422, "GEMINI_NOT_CONFIGURED", "Cấu hình GEMINI_API_KEY trên server để nhận dạng lại bằng Gemini")
    if provider == "google" and not google_ready():
        fail(
            422,
            "GOOGLE_NOT_CONFIGURED",
            "Upload JSON Google hợp lệ trong Cấu hình giọng nói trước khi nhận dạng lại",
        )
    existing = db.scalar(select(ReviewJob).where(ReviewJob.attempt_id == key, ReviewJob.status == "PENDING"))
    if existing:
        if existing.policy["provider"] != provider:
            fail(409, "REVIEW_PENDING", "Đang có yêu cầu nhận dạng lại bằng nhà cung cấp khác. Chờ hoàn tất trước khi đổi.")
        return data(existing, "status")
    audio = db.scalar(
        select(Upload).where(Upload.attempt_id == key, Upload.kind == "AUDIO", Upload.status == "COMPLETED")
    )
    if not audio:
        fail(409, "AUDIO_NOT_READY", "Chưa có audio gốc đã tải lên hoàn tất")
    previous = db.scalar(
        select(ReviewJob)
        .where(ReviewJob.attempt_id == key, ReviewJob.status == "COMPLETED")
        .order_by(ReviewJob.created_at.desc())
        .limit(1)
    )
    job = ReviewJob(
        attempt_id=key,
        requested_by=user.id,
        reason=body.reason,
        policy=policy(db)
        | {"provider": provider}
        | ({"model": settings().gemini_stt_model, "preprocessing": "off"} if provider == "gemini" else {}),
        original={
            "transcript": previous.result["transcript"] if previous else attempt.transcript,
            "stt_confidence": previous.result["stt_confidence"] if previous else attempt.stt_confidence,
            "submitted_transcript": attempt.transcript,
            "assessment": attempt.assessment,
            "audio_id": audio.id,
            "audio_sha256": audio.sha256,
        },
    )
    db.add(job)
    session.status, session.final_score = "REVIEW_REQUIRED", None
    db.flush()
    db.add(
        Audit(
            user_id=user.id,
            event=provider.upper() + "_REVIEW_REQUESTED",
            details={"job_id": job.id, "attempt_id": key, "reason": body.reason},
        )
    )
    db.commit()
    return data(job, "status")


@router.put("/users/{key}/role")
def change_role(key: str, body: s.RoleIn, db: Session = Depends(get_db), user=Depends(admin)):
    admins = db.scalars(
        select(User).where(User.role.in_(("SYSTEM_ADMIN", "ADMIN")), User.status == "ACTIVE").order_by(User.id).with_for_update()
    ).all()
    row = by_id(db, User, key, lock=True)
    if row.role in {"SYSTEM_ADMIN", "ADMIN"} and body.role not in {"SYSTEM_ADMIN", "ADMIN"} and row.status == "ACTIVE" and len(admins) <= 1:
        fail(409, "LAST_ADMIN", "Cần giữ ít nhất một quản trị viên đang hoạt động")
    before = row.role
    row.role = body.role
    db.add(
        Audit(
            user_id=user.id,
            event="USER_ROLE_CHANGED",
            details={"user_id": row.id, "before": before, "after": body.role},
        )
    )
    db.commit()
    return public_user(row)


@router.get("/courses/{course_id}/students")
def course_students(course_id: str, detail: bool = False, db: Session = Depends(get_db), user=Depends(staff)):
    course_access(db, course_id, user)
    direct_ids = set(
        db.scalars(select(CourseEnrollment.student_id).where(CourseEnrollment.course_id == course_id))
    )
    sections = db.scalars(select(Section).where(Section.course_id == course_id)).all()
    section_map = {s.id: s.code for s in sections}
    section_enrollments = (
        db.scalars(select(ExamEnrollment).where(ExamEnrollment.section_id.in_(list(section_map.keys())))).all()
        if section_map
        else []
    )
    student_section = {se.student_id: section_map.get(se.section_id, "") for se in section_enrollments}
    all_student_ids = direct_ids | set(student_section.keys())

    if not detail:
        return list(all_student_ids if section_map else direct_ids)

    users = db.scalars(select(User).where(User.id.in_(list(all_student_ids)))).all() if all_student_ids else []
    return [
        {
            "id": u.id,
            "username": u.username,
            "name": u.name,
            "status": u.status,
            "section_code": student_section.get(u.id, "Mặc định"),
        }
        for u in sorted(users, key=lambda x: x.username or "")
    ]


@router.post("/courses/{course_id}/students")
def enroll_students(course_id: str, body: s.AssignIn, db: Session = Depends(get_db), user=Depends(admin)):
    course_access(db, course_id, user)
    by_id(db, Course, course_id, lock=True)
    for key in set(body.student_ids):
        learner = by_id(db, User, key)
        if learner.status != "ACTIVE":
            fail(422, "INACTIVE_USER", "Tài khoản không hoạt động")
        if not db.scalar(
            select(CourseEnrollment.id).where(
                CourseEnrollment.course_id == course_id, CourseEnrollment.student_id == key
            )
        ):
            db.add(CourseEnrollment(course_id=course_id, student_id=key))
    db.add(
        Audit(
            user_id=user.id,
            event="COURSE_ENROLLED",
            details={"course_id": course_id, "user_ids": body.student_ids},
        )
    )
    db.commit()
    return {"ok": True}


@router.delete("/courses/{course_id}/students/{student_id}")
def unenroll_student(course_id: str, student_id: str, db: Session = Depends(get_db), user=Depends(admin)):
    course_access(db, course_id, user)
    row = db.scalar(
        select(CourseEnrollment).where(
            CourseEnrollment.course_id == course_id, CourseEnrollment.student_id == student_id
        )
    )
    if row:
        db.delete(row)
    db.add(
        Audit(
            user_id=user.id,
            event="COURSE_UNENROLLED",
            details={"course_id": course_id, "user_id": student_id},
        )
    )
    db.commit()
    return {"ok": True}


@router.put("/exams/{key}/attempt-policy")
def update_attempt_policy(key: str, body: s.AttemptPolicyIn, db: Session = Depends(get_db), user=Depends(admin)):
    exam = by_id(db, Exam, key, lock=True)
    previous = exam.max_attempts
    exam.max_attempts = body.max_attempts
    db.add(Audit(user_id=user.id, event="EXAM_ATTEMPT_POLICY", details={
        "exam_id": key, "previous": previous, "max_attempts": body.max_attempts,
    }))
    db.commit()
    return data(exam, "max_attempts")


@router.post("/results/{key}/retake")
def grant_retake(key: str, body: s.RetakeIn, db: Session = Depends(get_db), user=Depends(admin)):
    session = by_id(db, ExamSession, key)
    exam = by_id(db, Exam, session.exam_id, lock=True)
    db.refresh(session)
    if session.deleted_at is not None:
        fail(404, "NOT_FOUND", "Lần thi đã bị xóa")
    if exam.max_attempts is None:
        fail(409, "UNLIMITED", "Bài thi đã cho làm lại không giới hạn")
    state = allowance(db, exam, session.student_id)
    assignment = db.scalar(select(Assignment).where(
        Assignment.exam_id == exam.id, Assignment.student_id == session.student_id,
    ))
    if not assignment:
        assignment = Assignment(exam_id=exam.id, student_id=session.student_id, extra_attempts=0)
        db.add(assignment)
    assignment.extra_attempts = max(
        assignment.extra_attempts + body.additional_attempts,
        state["attempt_count"] + body.additional_attempts - exam.max_attempts,
    )
    db.add(Audit(user_id=user.id, event="RETAKE_GRANTED", details={
        "exam_id": exam.id, "student_id": session.student_id,
        "additional_attempts": body.additional_attempts, "extra_attempts": assignment.extra_attempts,
    }))
    db.commit()
    return allowance(db, exam, session.student_id)


@router.delete("/results/{key}")
def delete_result(key: str, db: Session = Depends(get_db), user=Depends(admin)):
    session = by_id(db, ExamSession, key)
    # Share the exam lock with session creation and allowance changes.
    by_id(db, Exam, session.exam_id, lock=True)
    ids = list(db.scalars(select(Attempt.id).where(Attempt.session_id == key)))
    try:
        # Worker lock order is review job -> attempt -> session. NOWAIT also avoids
        # conflicting with upload/finish transactions and returns a retryable UI error.
        db.scalars(select(ReviewJob).where(ReviewJob.attempt_id.in_(ids)).with_for_update(nowait=True)).all()
        db.scalars(select(Attempt).where(Attempt.id.in_(ids)).with_for_update(nowait=True)).all()
        session = db.scalar(select(ExamSession).where(ExamSession.id == key)
                            .execution_options(populate_existing=True).with_for_update(nowait=True))
        uploads = db.scalars(select(Upload).where(Upload.attempt_id.in_(ids)).with_for_update(nowait=True)).all()
    except OperationalError as exc:
        if getattr(exc.orig, "sqlstate", None) != "55P03":
            raise
        db.rollback()
        fail(409, "SESSION_BUSY", "Lần thi đang được xử lý. Đợi hoàn tất rồi xóa lại.")
    if session.deleted_at is not None:
        return {"ok": True}
    for upload in uploads:
        db.add(MediaCleanup(upload_id=upload.id, storage_key=upload.storage_key))
    db.execute(delete(ReviewJob).where(ReviewJob.attempt_id.in_(ids)))
    db.execute(delete(Upload).where(Upload.attempt_id.in_(ids)))
    db.execute(delete(Attempt).where(Attempt.id.in_(ids)))
    session.deleted_at, session.status, session.final_score = time.time(), "DELETED", None
    db.add(Audit(user_id=user.id, event="EXAM_SESSION_DELETED", details={
        "session_id": key, "exam_id": session.exam_id, "student_id": session.student_id,
        "attempt_number": session.attempt_number, "media_count": len(uploads),
    }))
    db.commit()
    return {"ok": True}


# ─────────────────────────────────────────────────────────────────
# Phase 3: Grading Review, Transcript Correction, Score Override & Re-evaluation
# ─────────────────────────────────────────────────────────────────

@router.post("/attempts/{key}/regrade-transcript")
def regrade_transcript(
    key: str,
    body: s.RegradeTranscriptIn,
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    attempt = by_id(db, Attempt, key, lock=True)
    session = by_id(db, ExamSession, attempt.session_id, lock=True)
    if session.deleted_at is not None:
        fail(404, "NOT_FOUND", "Lần thi đã bị xóa")
    exam = by_id(db, Exam, session.exam_id)
    course_access(db, exam.course_id, user)

    prev_transcript = attempt.transcript
    attempt.transcript = body.corrected_transcript

    if not attempt.finished_at:
        attempt.finished_at = time.time()

    try:
        assessment = grade_answer(
            db,
            exam,
            session,
            attempt,
            body.corrected_transcript,
            attempt.stt_confidence or 1.0,
        )
    except Exception as exc:
        fail(500, "REGRADE_FAILED", f"Lỗi chấm lại: {exc}")

    assessment["transcript_edited"] = True
    assessment["original_transcript"] = prev_transcript
    assessment["edit_reason"] = body.reason
    assessment["edited_by"] = user.name
    assessment["edited_at"] = time.time()

    attempt.assessment = assessment
    attempt.status = "GRADED"
    db.flush()

    finalize(db, session)

    db.add(
        Audit(
            user_id=user.id,
            event="TRANSCRIPT_REGRADED",
            details={
                "attempt_id": key,
                "session_id": session.id,
                "reason": body.reason,
                "score": assessment.get("score"),
            },
        )
    )
    db.commit()

    return data(attempt, "sequence", "question", "transcript", "stt_confidence", "assessment", "status") | {
        "assessment": assessment_view(attempt.assessment, exam)
    }


@router.post("/attempts/{key}/override")
def override_attempt_score(
    key: str,
    body: s.ScoreOverrideIn,
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    attempt = by_id(db, Attempt, key, lock=True)
    session = by_id(db, ExamSession, attempt.session_id, lock=True)
    if session.deleted_at is not None:
        fail(404, "NOT_FOUND", "Lần thi đã bị xóa")
    exam = by_id(db, Exam, session.exam_id)
    course_access(db, exam.course_id, user)

    assessment = dict(attempt.assessment or {})
    prev_score = assessment.get("score")
    assessment["score"] = round(body.score, 2)
    assessment["manual_override"] = True
    assessment["override_reason"] = body.reason
    assessment["override_by"] = user.name
    assessment["overridden_at"] = time.time()

    if body.criteria and "criteria" in assessment:
        override_map = {c.name: c for c in body.criteria}
        updated_criteria = []
        for crit in assessment.get("criteria", []):
            crit_copy = dict(crit)
            if crit_copy.get("name") in override_map:
                item = override_map[crit_copy["name"]]
                crit_copy["score"] = item.score
                if item.feedback:
                    crit_copy["feedback"] = item.feedback
            updated_criteria.append(crit_copy)
        assessment["criteria"] = updated_criteria

    attempt.assessment = assessment
    attempt.status = "GRADED"
    db.flush()

    # Recalculate session final score
    attempts = db.scalars(select(Attempt).where(Attempt.session_id == session.id)).all()
    scores = [a.assessment.get("score") for a in attempts if a.assessment]
    if scores and all(s is not None for s in scores):
        session.final_score = round(sum(scores) / len(scores), 2)

    db.add(
        Audit(
            user_id=user.id,
            event="ATTEMPT_SCORE_OVERRIDDEN",
            details={
                "attempt_id": key,
                "session_id": session.id,
                "prev_score": prev_score,
                "score": body.score,
                "reason": body.reason,
            },
        )
    )
    db.commit()
    return data(attempt, "sequence", "question", "transcript", "stt_confidence", "assessment", "status") | {
        "assessment": assessment_view(attempt.assessment, exam)
    }


@router.post("/results/{key}/approve")
def approve_result(
    key: str,
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    session = by_id(db, ExamSession, key, lock=True)
    if session.deleted_at is not None:
        fail(404, "NOT_FOUND", "Lần thi đã bị xóa")
    exam = by_id(db, Exam, session.exam_id)
    course_access(db, exam.course_id, user)

    attempts = db.scalars(select(Attempt).where(Attempt.session_id == key)).all()
    if not attempts:
        fail(409, "NO_ATTEMPTS", "Chưa có câu trả lời nào để duyệt")

    for a in attempts:
        if not a.assessment or a.assessment.get("score") is None:
            fail(409, "INCOMPLETE_GRADING", f"Câu {a.sequence} chưa hoàn tất chấm điểm")

    scores = [a.assessment.get("score") for a in attempts]
    session.final_score = round(sum(scores) / len(scores), 2)
    session.status = "COMPLETED"

    db.add(
        Audit(
            user_id=user.id,
            event="EXAM_SESSION_APPROVED",
            details={"session_id": key, "student_id": session.student_id, "final_score": session.final_score},
        )
    )
    db.commit()
    return {
        "id": session.id,
        "status": session.status,
        "final_score": session.final_score,
    }


@router.post("/attempts/{key}/request-re-eval")
def request_re_evaluation_admin(
    key: str,
    body: s.ReEvaluationIn,
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    attempt = by_id(db, Attempt, key)
    session = by_id(db, ExamSession, attempt.session_id)
    exam = by_id(db, Exam, session.exam_id)
    course_access(db, exam.course_id, user)

    by_id(db, User, body.teacher_id_2)

    original_score = attempt.assessment.get("score") if attempt.assessment else None

    row = ReEvaluation(
        attempt_id=key,
        teacher_id_1=user.id,
        teacher_id_2=body.teacher_id_2,
        reason=body.reason,
        reason_detail=body.reason_detail,
        blind_marking=body.blind_marking,
        score_1=original_score,
        status="PENDING",
    )
    db.add(row)
    db.add(
        Audit(
            user_id=user.id,
            event="REEVALUATION_REQUESTED",
            details={
                "attempt_id": key,
                "teacher_id_1": user.id,
                "teacher_id_2": body.teacher_id_2,
                "reason": body.reason,
            },
        )
    )
    db.commit()
    return {"id": row.id, "status": row.status}


@router.get("/re-evaluations")
def list_re_evaluations(
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    query = select(ReEvaluation).order_by(ReEvaluation.created_at.desc())
    if user.role == "TEACHER":
        query = query.where((ReEvaluation.teacher_id_1 == user.id) | (ReEvaluation.teacher_id_2 == user.id))
    rows = db.scalars(query).all()
    results = []
    for r in rows:
        att = db.get(Attempt, r.attempt_id)
        sess = db.get(ExamSession, att.session_id) if att else None
        student = db.get(User, sess.student_id) if sess else None
        t1 = db.get(User, r.teacher_id_1)
        t2 = db.get(User, r.teacher_id_2)
        is_blind = r.blind_marking and user.id == r.teacher_id_2 and r.status != "COMPLETED"
        results.append({
            "id": r.id,
            "attempt_id": r.attempt_id,
            "session_id": sess.id if sess else None,
            "student_name": student.name if student else "N/A",
            "teacher_1_name": t1.name if t1 else "Giảng viên 1",
            "teacher_2_name": t2.name if t2 else "Giảng viên 2",
            "reason": r.reason,
            "reason_detail": r.reason_detail,
            "status": r.status,
            "score_1": None if is_blind else r.score_1,
            "score_2": r.score_2,
            "final_score": r.final_score,
            "blind_marking": r.blind_marking,
            "created_at": r.created_at,
            "completed_at": r.completed_at,
        })
    return results


@router.post("/re-evaluations/{key}/submit")
def submit_re_evaluation(
    key: str,
    body: s.ReEvaluationSubmitIn,
    db: Session = Depends(get_db),
    user=Depends(staff),
):
    row = by_id(db, ReEvaluation, key, lock=True)
    if user.role == "TEACHER" and row.teacher_id_2 != user.id:
        fail(403, "FORBIDDEN", "Bạn không phải giảng viên được phân công chấm thẩm định này")

    row.score_2 = round(body.score_2, 2)
    if row.score_1 is not None:
        row.final_score = round((row.score_1 + row.score_2) / 2, 2)
    else:
        row.final_score = row.score_2
    row.status = "COMPLETED"
    row.completed_at = time.time()

    att = db.get(Attempt, row.attempt_id)
    if att:
        assessment = dict(att.assessment or {})
        assessment["score"] = row.final_score
        assessment["re_evaluated"] = True
        assessment["re_eval_id"] = row.id
        att.assessment = assessment
        sess = db.get(ExamSession, att.session_id)
        if sess:
            attempts = db.scalars(select(Attempt).where(Attempt.session_id == sess.id)).all()
            scores = [a.assessment.get("score") for a in attempts if a.assessment]
            if scores and all(s is not None for s in scores):
                sess.final_score = round(sum(scores) / len(scores), 2)

    db.add(
        Audit(
            user_id=user.id,
            event="REEVALUATION_COMPLETED",
            details={"re_eval_id": row.id, "score_2": row.score_2, "final_score": row.final_score},
        )
    )
    db.commit()
    return {
        "id": row.id,
        "status": row.status,
        "score_2": row.score_2,
        "final_score": row.final_score,
    }
