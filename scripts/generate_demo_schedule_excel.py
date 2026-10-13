import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_demo_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Phan_Ca_Thi_Sinh_Vien"

    headers = [
        "Ca_Thi",
        "Ngay_Thi",
        "Gio_Bat_Dau",
        "Gio_Ket_Thuc",
        "Phong_Thi",
        "Giang_Vien_Email",
        "MSSV",
        "Ho_Ten",
        "Lop_Hoc"
    ]
    ws.append(headers)

    # Styling header
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    sample_data = [
        # Ca 1 - Phong AL-L401 (Giang vien: teacher1@fpt.edu.vn)
        [1, "2026-10-25", "08:00", "09:30", "AL-L401", "teacher1@fpt.edu.vn", "SE170101", "Nguyễn Hoàng Long", "SE1701"],
        [1, "2026-10-25", "08:00", "09:30", "AL-L401", "teacher1@fpt.edu.vn", "SE170102", "Trần Thị Mai Anh", "SE1701"],
        [1, "2026-10-25", "08:00", "09:30", "AL-L401", "teacher1@fpt.edu.vn", "SE170103", "Lê Minh Tuấn", "SE1701"],
        [1, "2026-10-25", "08:00", "09:30", "AL-L401", "teacher1@fpt.edu.vn", "SE170104", "Phạm Quỳnh Nga", "SE1701"],

        # Ca 1 - Phong AL-L402 (Giang vien: teacher_coicham@fpt.edu.vn)
        [1, "2026-10-25", "08:00", "09:30", "AL-L402", "teacher_coicham@fpt.edu.vn", "SE170201", "Đỗ Quang Dũng", "SE1702"],
        [1, "2026-10-25", "08:00", "09:30", "AL-L402", "teacher_coicham@fpt.edu.vn", "SE170202", "Hoàng Bảo Yến", "SE1702"],
        [1, "2026-10-25", "08:00", "09:30", "AL-L402", "teacher_coicham@fpt.edu.vn", "SE170203", "Vũ Trọng Hưng", "SE1702"],
        [1, "2026-10-25", "08:00", "09:30", "AL-L402", "teacher_coicham@fpt.edu.vn", "SE170204", "Đặng Phương Thảo", "SE1702"],

        # Ca 2 - Phong AL-L401 (Giang vien: teacher1@fpt.edu.vn)
        [2, "2026-10-25", "10:00", "11:30", "AL-L401", "teacher1@fpt.edu.vn", "SE170105", "Bùi Văn Hậu", "SE1701"],
        [2, "2026-10-25", "10:00", "11:30", "AL-L401", "teacher1@fpt.edu.vn", "SE170106", "Dương Thu Hằng", "SE1701"],
        [2, "2026-10-25", "10:00", "11:30", "AL-L401", "teacher1@fpt.edu.vn", "SE170107", "Ngô Quốc Bảo", "SE1701"],

        # Ca 2 - Phong AL-L402 (Giang vien: teacher_coicham@fpt.edu.vn)
        [2, "2026-10-25", "10:00", "11:30", "AL-L402", "teacher_coicham@fpt.edu.vn", "SE170205", "Trịnh Khánh Linh", "SE1702"],
        [2, "2026-10-25", "10:00", "11:30", "AL-L402", "teacher_coicham@fpt.edu.vn", "SE170206", "Phan Gia Khiêm", "SE1702"],
        [2, "2026-10-25", "10:00", "11:30", "AL-L402", "teacher_coicham@fpt.edu.vn", "SE170207", "Nguyễn Thảo My", "SE1702"]
    ]

    for row_idx, r in enumerate(sample_data, start=2):
        ws.append(r)
        for col_idx in range(1, len(r) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = Font(name="Calibri", size=10)
            c.border = thin_border
            if col_idx in (1, 2, 3, 4, 5, 7, 9):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    p1 = os.path.join(base_dir, "sample_data", "demo_import_ca_thi_sinh_vien.xlsx")
    p2 = os.path.join(base_dir, "apps", "staff-portal", "public", "demo_import_ca_thi_sinh_vien.xlsx")

    os.makedirs(os.path.dirname(p1), exist_ok=True)
    os.makedirs(os.path.dirname(p2), exist_ok=True)

    wb.save(p1)
    wb.save(p2)
    print("Demo files successfully generated at:")
    print("1.", p1)
    print("2.", p2)

if __name__ == "__main__":
    create_demo_excel()
